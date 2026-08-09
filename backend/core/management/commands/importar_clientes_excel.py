"""
Importa el listado real de clientes de Panaderia Macarron desde el Excel
de rutas que usan hoy (hojas "CLIENTES" y "CLIENTES DELIVERY"). Uso:

    python manage.py importar_clientes_excel "C:\\ruta\\al\\archivo.xlsx" --empresa "Panadería Macarrón"

Es idempotente: si vuelves a correrlo, no duplica clientes ya importados
(se identifican por nombre + direccion dentro de la misma empresa) pero
SI actualiza su tipo_cliente (normal/delivery) por si cambio en el Excel.

Clasificacion: un cliente que aparece en la hoja "CLIENTES" (el reparto
comercial/furgon) queda como NORMAL aunque tambien aparezca en
"CLIENTES DELIVERY"; si aparece SOLO en "CLIENTES DELIVERY" queda como
DELIVERY (entrega a domicilio de particulares).
"""
import openpyxl
from django.core.management.base import BaseCommand, CommandError

from core.models import Empresa, Cliente

HOJA_NORMAL = "CLIENTES "
HOJA_DELIVERY = "CLIENTES DELIVERY"


class Command(BaseCommand):
    help = "Importa/actualiza clientes reales desde el Excel de rutas de la panaderia."

    def add_arguments(self, parser):
        parser.add_argument("archivo", type=str, help="Ruta al archivo .xlsx")
        parser.add_argument("--empresa", type=str, default="Panadería Macarrón")
        parser.add_argument("--ciudad", type=str, default="Santiago")

    def _leer_hoja(self, wb, nombre_hoja):
        if nombre_hoja not in wb.sheetnames:
            self.stdout.write(self.style.WARNING(f"Hoja '{nombre_hoja}' no encontrada, se omite."))
            return {}
        ws = wb[nombre_hoja]
        filas = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            fantasia, direccion, comuna = row[1], row[2], row[3]
            if not fantasia or not direccion:
                continue
            nombre = str(fantasia).strip()
            direccion = str(direccion).strip()
            clave = (nombre.upper(), direccion.upper())
            filas[clave] = {
                "nombre": nombre.title(),
                "direccion": direccion.title(),
                "comuna": str(comuna).strip().title() if comuna else "",
            }
        return filas

    def handle(self, *args, **options):
        ruta_archivo = options["archivo"]
        nombre_empresa = options["empresa"]
        ciudad_default = options["ciudad"]

        try:
            empresa = Empresa.objects.get(nombre=nombre_empresa)
        except Empresa.DoesNotExist:
            raise CommandError(
                f"No existe una Empresa llamada '{nombre_empresa}'. "
                "Revisa el nombre exacto o crea la empresa primero."
            )

        try:
            wb = openpyxl.load_workbook(ruta_archivo, data_only=True)
        except FileNotFoundError:
            raise CommandError(f"No se encontró el archivo: {ruta_archivo}")

        normales = self._leer_hoja(wb, HOJA_NORMAL)
        delivery = self._leer_hoja(wb, HOJA_DELIVERY)

        # normales tiene prioridad: si esta en ambas hojas, queda NORMAL.
        combinados = {}
        for clave, datos in delivery.items():
            combinados[clave] = {**datos, "tipo_cliente": "DELIVERY"}
        for clave, datos in normales.items():
            combinados[clave] = {**datos, "tipo_cliente": "NORMAL"}

        creados = 0
        actualizados = 0
        sin_cambios = 0

        for datos in combinados.values():
            cliente = Cliente.objects.filter(
                empresa=empresa, nombre__iexact=datos["nombre"], direccion__iexact=datos["direccion"]
            ).first()

            if cliente is None:
                Cliente.objects.create(
                    empresa=empresa,
                    nombre=datos["nombre"],
                    direccion=datos["direccion"],
                    comuna=datos["comuna"],
                    ciudad=ciudad_default,
                    tipo_cliente=datos["tipo_cliente"],
                    activo=True,
                )
                creados += 1
            elif cliente.tipo_cliente != datos["tipo_cliente"]:
                cliente.tipo_cliente = datos["tipo_cliente"]
                cliente.save(update_fields=["tipo_cliente"])
                actualizados += 1
            else:
                sin_cambios += 1

        self.stdout.write(self.style.SUCCESS(
            f"Listo. Creados: {creados}. Tipo actualizado: {actualizados}. Sin cambios: {sin_cambios}."
        ))
        self.stdout.write(
            "Nota: los clientes nuevos NO quedan geocodificados automaticamente. Corre por separado:\n"
            "  python manage.py regeocodificar_clientes"
        )
